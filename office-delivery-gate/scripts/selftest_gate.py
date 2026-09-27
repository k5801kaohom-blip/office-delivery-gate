#!/usr/bin/env python3
"""Prove the delivery gate actually works by breaking a file on purpose.

A validator that has never failed anything is not evidence of safety. This script
injects a real, historically-observed defect into a copy of a clean Office file and
asserts the whole chain behaves: validator catches it, repairer fixes it,
re-validation passes, and the repair preserves the document's content.

Exit codes
    0   the gate behaved correctly on every assertion
    1   an assertion failed - do NOT trust the gate until this is fixed
    2   environment error (missing helper, unusable input)
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

_HERE = Path(__file__).resolve().parent
_CANDIDATES = [
    Path("/home/ubuntu/skills/ooxml-integrity-check/scripts"),
    _HERE.parent.parent / "ooxml-integrity-check" / "scripts",
]
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def find_helper(name: str) -> Path:
    for base in _CANDIDATES:
        candidate = base / name
        if candidate.is_file():
            return candidate
    raise SystemExit(f"[selftest] cannot locate {name}; install ooxml-integrity-check nearby")


def run(cmd: list[str]) -> tuple[int, str]:
    proc = subprocess.run(cmd, capture_output=True, text=True)
    return proc.returncode, (proc.stdout + proc.stderr).strip()


def inject_duplicate_latin(src: Path, dst: Path) -> int:
    """Reproduce the generator bug that produced unopenable decks.

    python-pptx's `run.font.name = "..."` already creates <a:latin> inside <a:rPr>.
    Appending a second <a:latin> for the East Asian face yields
    <a:latin/><a:latin/><a:ea/><a:cs/>, which PowerPoint reports as corrupt while
    LibreOffice still renders it perfectly.
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation(str(src))
    touched = 0
    for index, slide in enumerate(prs.slides, start=1):
        box = slide.shapes.add_textbox(Inches(0.3), Inches(0.2), Inches(3.5), Inches(0.4))
        frame = box.text_frame
        frame.text = f"gate selftest {index}"
        run_ = frame.paragraphs[0].runs[0]
        run_.font.size = Pt(12)
        run_.font.name = "Microsoft JhengHei"          # creates <a:latin>
        rpr = run_._r.get_or_add_rPr()
        for tag in ("latin", "ea", "cs"):              # BUG: a second <a:latin>
            rpr.append(rpr.makeelement(f"{{{A_NS}}}{tag}", {"typeface": "Microsoft JhengHei"}))
        touched += 1
    prs.save(str(dst))
    return touched


def content_counts(path: Path) -> dict:
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        return {
            "slides": len([n for n in names if n.startswith("ppt/slides/slide")]),
            "media": len([n for n in names if n.startswith("ppt/media/")]),
        }


def main() -> int:
    ap = argparse.ArgumentParser(description="Self-test the Office delivery gate.")
    ap.add_argument("sample", help="a clean .pptx to use as the test subject")
    ap.add_argument("--workdir", default=None, help="scratch directory (default: ./gate_selftest)")
    args = ap.parse_args()

    sample = Path(args.sample)
    if not sample.is_file():
        print(f"[selftest] sample not found: {sample}", file=sys.stderr)
        return 2

    workdir = Path(args.workdir or "gate_selftest")
    if workdir.exists():
        shutil.rmtree(workdir)
    workdir.mkdir(parents=True)

    try:
        validate = find_helper("validate_ooxml.py")
        repair = find_helper("repair_ooxml.py")
    except SystemExit as exc:
        print(exc, file=sys.stderr)
        return 2

    gate = _HERE / "delivery_gate.py"
    failures: list[str] = []

    def check(label: str, ok: bool, detail: str = "") -> None:
        print(f"  {'PASS' if ok else 'FAIL'}  {label}" + (f"   {detail}" if detail else ""))
        if not ok:
            failures.append(label)

    print("[selftest] 1/5 clean input must clear the gate as-is")
    rc, _ = run([sys.executable, str(validate), str(sample), "--quiet"])
    check("clean file passes validation", rc == 0, f"exit={rc}")

    print("[selftest] 2/5 inject the real-world defect")
    broken = workdir / "injected_defect.pptx"
    try:
        touched = inject_duplicate_latin(sample, broken)
    except Exception as exc:  # noqa: BLE001 - surface any injection problem plainly
        print(f"[selftest] injection failed: {exc}", file=sys.stderr)
        return 2
    check("defect injected", touched > 0 and broken.is_file(), f"{touched} slide(s)")

    print("[selftest] 3/5 validator must catch it")
    rc_broken, out_broken = run([sys.executable, str(validate), str(broken), "--quiet"])
    caught = rc_broken == 1 and "duplicate singleton children" in out_broken
    check("validator rejects the defective file", caught, f"exit={rc_broken}")
    if not caught:
        print(out_broken)

    print("[selftest] 4/5 gate with --repair must recover it")
    report = workdir / "gate_report.json"
    rc_gate, out_gate = run([
        sys.executable, str(gate), str(broken),
        "--repair", "--workdir", str(workdir / "out"), "--json", str(report),
    ])
    check("gate clears the repairable file", rc_gate == 0, f"exit={rc_gate}")
    if not report.is_file():
        check("gate report written", False)
        return 1
    payload = json.loads(report.read_text(encoding="utf-8"))
    check("gate verdict is SAFE", payload.get("verdict") == "SAFE")

    print("[selftest] 5/5 repair must preserve content")
    repaired = Path(payload["results"][0].get("repaired_path", ""))
    if repaired.is_file():
        before, after = content_counts(broken), content_counts(repaired)
        check("slides preserved", before["slides"] == after["slides"],
              f"{before['slides']} -> {after['slides']}")
        check("media preserved", before["media"] == after["media"],
              f"{before['media']} -> {after['media']}")
    else:
        check("repaired file exists", False)

    print()
    if failures:
        print(f"[selftest] {len(failures)} assertion(s) failed - the gate is NOT trustworthy yet:")
        for f in failures:
            print(f"          - {f}")
        return 1
    print("[selftest] all assertions passed - the gate catches and repairs a real defect.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())